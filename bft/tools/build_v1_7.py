"""Builds Battery_Fuse_Block_Cost_Estimator_v1.7.xlsx on top of the v1.6 build.

v1.7 adds:
  * Ambient temperature (Estimator B15; underhood 105 C, interior 85 C) and a max busbar
    temperature (Assumptions B69). The Busbar Sizer's allowed rise is max temp minus ambient,
    and the rating is lowered for copper's higher resistance at the max temperature.
    The Estimator shows the estimated bus temperature at load (it still sizes the bus by
    current density).
  * Busbar alloy conductivity, % IACS (Assumptions B72). The ampacity and the design current
    density scale with its square root.
  * Output-side cover skirt drop (Assumptions B77, 20 mm): the cover over the output wires
    reaches below the eyelet's conductive area. A check shows how far that skirt reaches
    below the base floor.
  * Busbar Sizer chart rebuilt (it rendered empty) and MINIFS replaced by a helper column so
    the sheet works in any Excel version.
  * Cached values are written into the file, so it shows numbers and the chart even in
    Excel Protected View, before Excel recalculates.

The web tool (bft/web/bft_estimator_v1.7.html) must match these formulas.
Run: python bft/tools/build_v1_7.py
"""
import copy
import os
import subprocess
import sys
import tempfile
import zipfile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.chart import Reference, ScatterChart, Series

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_v1_6 as v6  # noqa: E402
from build_v1_4 import N0, N1, N2, calc, f, head, inp  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "Battery_Fuse_Block_Cost_Estimator_v1.7.xlsx")
A = "Assumptions!"
K_OLD = "Assumptions!$B$58*Assumptions!$B$63"
K_NEW = "Assumptions!$B$58*Assumptions!$B$63*Assumptions!$B$73"


def style(dst, src):
    dst.font = copy.copy(src.font)
    dst.fill = copy.copy(src.fill)
    dst.border = copy.copy(src.border)
    dst.number_format = src.number_format
    dst.alignment = copy.copy(src.alignment)


def edit(wb):
    est, asm, bs = wb["Estimator"], wb["Assumptions"], wb["Busbar Sizer"]

    # ---------------------------------------------------------------- Assumptions
    head(asm, 68, "Busbar temperature")
    rows = [(69, "Max busbar temperature", 150, N0, "C",
             "Placeholder. Set from the plastic, plating and fuse terminal limits"),
            (70, "Copper resistance temperature coefficient", 0.00393, "0.00000", "1/C", "Annealed copper"),
            (71, "Reference ambient for k", 35, N0, "C", "Ambient of the DIN 43671 ratings k was fitted to"),
            (72, "Busbar alloy conductivity", 100, N0, "% IACS",
             "100 = pure copper (C11000 about 101). Lower for alloys, e.g. CuFe2P (C19400) about 65")]
    for r, label, v, fmt, u, note in rows:
        asm.cell(r, 1, label).font = f()
        inp(asm.cell(r, 2), v, fmt)
        asm.cell(r, 3, u).font = f()
        asm.cell(r, 4, note).font = f("595959", size=9)
    for r, label, fm, u, note in (
            (73, "Rating factor (alloy and max temperature)",
             "=SQRT(B72/100*(1+B70*(B71+B62-20))/(1+B70*(B69-20)))", "ratio",
             "Calculated. Rating scales with the square root of conductivity, which drops as the bar heats"),
            (74, "Design current density used for the bus", "=B14*SQRT(B72/100)", "A/mm2",
             "Calculated. B14 is for 100% IACS copper, scaled by the square root of the alloy conductivity")):
        asm.cell(r, 1, label).font = f()
        calc(asm.cell(r, 2), fm, "0.000")
        asm.cell(r, 3, u).font = f()
        asm.cell(r, 4, note).font = f("595959", size=9)
    asm["A14"] = "Copper design current density (100% IACS)"
    head(asm, 76, "Cover isolation at the outputs")
    asm["A77"] = "Output-side cover skirt drop"
    asm["A77"].font = f()
    inp(asm["B77"], 20, N1)
    asm["C77"] = "mm"
    asm["C77"].font = f()
    asm["D77"] = ("Placeholder. Extra skirt length where the output wires leave, so the cover reaches below "
                  "the eyelet's conductive area")
    asm["D77"].font = f("595959", size=9)

    # ---------------------------------------------------------------- Estimator
    est["A3"] = ("Revision v1.7: ambient and max bus temperature, output-side cover skirt drop, Busbar Sizer "
                 "chart fix. Builds on v1.6.")
    for c, src in (("A15", "A5"), ("B15", "B5"), ("C15", "C5")):
        style(est[c], est[src])
    est["A15"] = "Ambient temperature"
    est["B15"] = 105
    est["B15"].number_format = N0
    est["C15"] = "C (underhood 105, interior 85)"
    for c, src in (("A31", "A30"), ("B31", "B30"), ("C31", "C30")):
        style(est[c], est[src])
    est["A31"] = "Estimated bus temperature at load"
    est["B31"] = (f"=B15+{A}$B$62*(B6/({K_NEW}*(B28*B7)^{A}$B$59*(2*(B28+B7))^{A}$B$60))^(1/{A}$B$61)")
    est["B31"].number_format = N1
    est["C31"] = "C"
    style(est["D15"], est["D14"])
    style(est["E15"], est["E14"])
    est["D15"] = "Bus temperature at load"
    est["E15"] = f'=TEXT(B31,"0.0")&" C, limit "&TEXT({A}B69,"0")&" C"&IF(B31>{A}B69,"  OVER","")'
    est["A63"] = "Skirt volume (incl. output-side drop)"
    est["B63"] = ("=2*(E65+E66-2*Assumptions!$B$8)*Assumptions!$B$8*Assumptions!$B$9"
                  "+Assumptions!$B$77*Assumptions!$B$8*(IF(Layout!B52>0,E66+2*E64,0)+IF(Layout!B51>0,E65+2*E64,0))")
    style(est["D67"], est["D66"])
    style(est["E67"], est["E66"])
    est["D67"] = "Output skirt below base floor (mm)"
    est["E67"] = "=MAX(0,Assumptions!$B$9+Assumptions!$B$77-(Assumptions!$B$4+B37))"

    # bus sizing by current density uses the alloy-adjusted density
    est["B26"] = "=B6/(Assumptions!$B$74*B7)"
    ly = wb["Layout"]
    for c in ("B42", "C42"):
        ly[c] = ly[c].value.replace("/(Assumptions!$B$14*", "/(Assumptions!$B$74*")
        assert "$B$74" in ly[c].value

    # ---------------------------------------------------------------- Busbar Sizer
    bs["A6"] = "Allowed temperature rise (max bus temp - ambient)"
    bs["B6"] = f"={A}$B$69-Estimator!$B$15"
    bs["B6"].font = f("008000")
    bs["C6"] = "C. Ambient is Estimator B15, max temp is Assumptions B69"
    for row in bs.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and K_OLD in c.value:
                c.value = c.value.replace(K_OLD, K_NEW)
    # MINIFS needs Excel 2019+, and AGGREGATE over an array fails in LibreOffice, so use a helper column
    style(bs["I25"], bs["H25"])
    bs["I25"] = "Cost if it fits ($)"
    bs.column_dimensions["I"].width = 12
    for r in range(26, 38):
        style(bs.cell(r, 9), bs.cell(r, 7))
        bs.cell(r, 9).value = f'=IF(H{r}="Yes",G{r},"")'
    bs["B39"] = '=IF(COUNT(I26:I37)=0,"None fit",INDEX(A26:A37,MATCH(MIN(I26:I37),I26:I37,0)))'
    bs["B40"] = '=IF(COUNT(I26:I37)=0,"-",MIN(I26:I37))'
    bs.cell(2, 1).value = ("Quick copper cost for a known bar, and the cheapest width-to-thickness ratio at a fixed "
                           "current. Ampacity model: Assumptions rows 57 to 63 and 69 to 73.")

    bs._charts = []
    ch = ScatterChart()
    ch.title = "Copper cost vs width : thickness"
    ch.scatterStyle = "lineMarker"
    ch.style = 2
    ch.legend = None
    ch.height, ch.width = 9, 17
    ch.x_axis.title = "Width : thickness"
    ch.y_axis.title = "Copper cost ($)"
    ch.x_axis.axPos = "b"
    ch.y_axis.axPos = "l"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.number_format = "$0.00"
    ch.x_axis.number_format = "0"
    ch.y_axis.scaling.min = 0
    ch.x_axis.scaling.min = 0
    s = Series(Reference(bs, min_col=7, min_row=26, max_row=37),
               Reference(bs, min_col=1, min_row=26, max_row=37), title="Cost ($)")
    s.graphicalProperties.line.solidFill = "2F6FA8"
    s.graphicalProperties.line.width = 28575
    s.marker.symbol = "circle"
    s.marker.size = 7
    s.marker.graphicalProperties.solidFill = "2F6FA8"
    s.marker.graphicalProperties.line.solidFill = "2F6FA8"
    s.smooth = False
    ch.series.append(s)
    bs.add_chart(ch, "J4")


def add_cached_values(path):
    """Recalculate in LibreOffice and write the results into the file as cached values."""
    d = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", d, path],
                   capture_output=True, timeout=300, check=True)
    calc_wb = load_workbook(os.path.join(d, os.path.basename(path)), data_only=True)
    wb = load_workbook(path)  # sheet order and names map sheetN.xml below
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cns = "http://schemas.openxmlformats.org/drawingml/2006/chart"

    zin = zipfile.ZipFile(path)
    files = {n: zin.read(n) for n in zin.namelist()}
    zin.close()

    # workbook.xml.rels maps sheets to files; openpyxl writes them as sheet1..N in order
    for i, ws in enumerate(wb.worksheets, start=1):
        name = f"xl/worksheets/sheet{i}.xml"
        root = etree.fromstring(files[name])
        vals = calc_wb[ws.title]
        for c in root.iterfind(".//m:c", ns):
            fnode = c.find("m:f", ns)
            if fnode is None:
                continue
            v = vals[c.get("r")].value
            if v is None:
                continue
            old = c.find("m:v", ns)
            if old is not None:
                c.remove(old)
            vnode = etree.SubElement(c, f"{{{ns['m']}}}v")
            if isinstance(v, bool):
                c.set("t", "b")
                vnode.text = "1" if v else "0"
            elif isinstance(v, (int, float)):
                c.attrib.pop("t", None)
                vnode.text = repr(float(v)) if isinstance(v, float) else str(v)
            elif isinstance(v, str) and v.startswith("#"):
                c.set("t", "e")
                vnode.text = v
            else:
                c.set("t", "str")
                vnode.text = str(v)
        files[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # chart caches so the chart draws before Excel recalculates
    for n in [n for n in files if n.startswith("xl/charts/chart") and n.endswith(".xml")]:
        root = etree.fromstring(files[n])
        for ref in root.iter(f"{{{cns}}}numRef"):
            fexpr = ref.find(f"{{{cns}}}f").text
            sheet, rng = fexpr.rsplit("!", 1)
            sheet = sheet.strip("'")
            a1, a2 = rng.replace("$", "").split(":")
            cells = calc_wb[sheet][f"{a1}:{a2}"]
            pts = [row[0].value for row in cells]
            for old in ref.findall(f"{{{cns}}}numCache"):
                ref.remove(old)
            cache = etree.SubElement(ref, f"{{{cns}}}numCache")
            etree.SubElement(cache, f"{{{cns}}}formatCode").text = "General"
            etree.SubElement(cache, f"{{{cns}}}ptCount").set("val", str(len(pts)))
            for idx, v in enumerate(pts):
                if isinstance(v, (int, float)):
                    pt = etree.SubElement(cache, f"{{{cns}}}pt")
                    pt.set("idx", str(idx))
                    etree.SubElement(pt, f"{{{cns}}}v").text = repr(float(v))
        files[n] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in files.items():
            zout.writestr(n, data)


def build(out=OUT, cache=True):
    tmp = os.path.join(tempfile.mkdtemp(), "v16.xlsx")
    v6.build(tmp)
    wb = load_workbook(tmp)
    edit(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    if cache:
        add_cached_values(out)
    return out


if __name__ == "__main__":
    print(build(*sys.argv[1:]))
