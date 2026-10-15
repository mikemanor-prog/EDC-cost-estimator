"""Builds Battery_Fuse_Block_Cost_Estimator_v1.8.xlsx on top of the v1.7 build.

v1.8 adds:
  * Bus sizing method (Assumptions B80). "Temperature limit" (default, Mike 2026-10-15) sizes the
    bus width so the bus reaches the max bus temperature (Assumptions B69) at the load and ambient.
    "Current density" keeps the v1.7 sizing (design A/mm2). The width is solved with 6 Newton
    steps in Layout rows 100 to 107 (the ampacity model has no closed form for width).
    The bolt width still sets a minimum. Zone branches stay sized by current density.
  * Copper scrap credit placeholder 50% (Assumptions B66), to be set with manufacturing.

The web tool (bft/web/bft_estimator_v1.8.html) must match these formulas.
Run: python bft/tools/build_v1_8.py
"""
import os
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_v1_7 as v7  # noqa: E402
from build_v1_4 import N1, calc, f, head, inp, unit  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "Battery_Fuse_Block_Cost_Estimator_v1.8.xlsx")
A = "Assumptions!"
NEWTON = 6


def edit(wb):
    est, asm, ly = wb["Estimator"], wb["Assumptions"], wb["Layout"]

    # ---------------------------------------------------------------- Assumptions
    asm["B66"] = 0.5
    asm["D66"] = ("Placeholder, to be set with manufacturing. Share of the copper price recovered for blank "
                  "scrap. Plated or mixed scrap usually returns less than clean bare copper")
    head(asm, 79, "Bus sizing")
    asm["A80"] = "Bus sizing method"
    asm["A80"].font = f()
    inp(asm["B80"], "Temperature limit")
    asm["D80"] = ("Temperature limit: the bus reaches the max bus temperature (B69) at the load and ambient. "
                  "Current density: design A/mm2 (B74). The bolt width is a minimum either way")
    asm["D80"].font = f("595959", size=9)
    d = DataValidation(type="list", formula1='"Temperature limit,Current density"', showErrorMessage=True)
    d.add("B80")
    asm.add_data_validation(d)

    # ---------------------------------------------------------------- Layout helper: width at the temperature limit
    head(ly, 99, "BUS WIDTH AT THE TEMPERATURE LIMIT (NEWTON SOLVE)", 4)
    k = v7.K_NEW
    ly["A100"] = "Load over k at the allowed rise"
    ly["A100"].font = f()
    calc(ly["B100"], f"=Estimator!$B$6/({k}*(MAX(1,{A}$B$69-Estimator!$B$15)/{A}$B$62)^{A}$B$61)", "0.000")
    unit(ly["C100"], "Allowed rise = max bus temp - ambient (at least 1 C)")
    ly["A101"] = "Start: width by current density"
    ly["A101"].font = f()
    calc(ly["B101"], f"=Estimator!$B$6/({A}$B$74*Estimator!$B$7)", N1)
    unit(ly["C101"], "mm")
    t, a, p = "Estimator!$B$7", f"{A}$B$59", f"{A}$B$60"
    for i in range(NEWTON):
        r = 102 + i
        w = f"B{r - 1}"
        ly.cell(r, 1, f"Newton step {i + 1}").font = f()
        calc(ly.cell(r, 2),
             f"=IF(Estimator!$B$6<=0,0,{w}*EXP(-({a}*LN({w}*{t})+{p}*LN(2*({w}+{t}))-LN($B$100))"
             f"/({a}+{p}*{w}/({w}+{t}))))", N1)
        unit(ly.cell(r, 3), "mm")
    last = 101 + NEWTON

    # ---------------------------------------------------------------- Estimator
    est["A3"] = ("Revision v1.8: bus sized by the temperature limit (or current density, Assumptions B80), "
                 "scrap credit placeholder 50%. Builds on v1.7.")
    est["B26"] = f'=IF({A}$B$80="Current density",Layout!$B$101,Layout!$B${last})'
    est["A26"] = "Bus width needed for the current (by the sizing method)"
    # the temperature-limit width puts B31 right at the limit, so flag OVER only beyond 0.05 C
    est["E15"] = f'=TEXT(B31,"0.0")&" C, limit "&TEXT({A}B69,"0")&" C"&IF(B31>{A}B69+0.05,"  OVER","")'
    est["B29"] = f'=IF(B26>=B27,IF({A}$B$80="Current density","Current load","Temperature limit"),"Bolt size")'


def build(out=OUT, cache=True):
    tmp = os.path.join(tempfile.mkdtemp(), "v17.xlsx")
    v7.build(tmp, cache=False)
    wb = load_workbook(tmp)
    edit(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    if cache:
        v7.add_cached_values(out)
    return out


if __name__ == "__main__":
    print(build(*sys.argv[1:]))
