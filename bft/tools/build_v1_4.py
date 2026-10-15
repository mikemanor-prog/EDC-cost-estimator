"""Builds Battery_Fuse_Block_Cost_Estimator_v1.4.xlsx from the v1.3 workbook.

v1.4 adds:
  * Keepout Zones sheet: two user-defined rectangular zones (bolted relay or other box)
    that add space to the base, with 0, 1 or 2 studs and an optional copper branch
    from the input bus.
  * Busbar Sizer sheet: cost of a known bar, plus a width-to-thickness ratio sweep at
    fixed current using an area + perimeter ampacity model.

The web tool (bft/web/bft_estimator_v1.4.html) must match these formulas.
Run: python bft/tools/build_v1_4.py
"""
import copy
import os
import sys

from openpyxl import load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "Battery_Fuse_Block_Cost_Estimator_v1.3.xlsx")
OUT = os.path.join(HERE, "..", "Battery_Fuse_Block_Cost_Estimator_v1.4.xlsx")

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3A4D")
SUB_FILL = PatternFill("solid", fgColor="DCE6EC")
INPUT_FILL = PatternFill("solid", fgColor="FFF9C4")
RESULT_FILL = PatternFill("solid", fgColor="E8F3E8")
THIN = Side(style="thin", color="B7BFCB")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
MONEY = "$0.00"
N0, N1, N2 = "0", "0.0", "0.00"

KZ = "'Keepout Zones'"
A = "Assumptions!"


def f(color=None, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def head(ws, row, text, span=3, col=1):
    ws.cell(row, col, text).font = f("FFFFFF", bold=True)
    for c in range(col, col + span):
        ws.cell(row, c).fill = HEAD_FILL


def sub(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font = f(bold=True)
        c.fill = SUB_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="left" if i == 0 else "center", wrap_text=True)


def inp(c, value, fmt=None):
    c.value = value
    c.font = f("0000FF")
    c.fill = INPUT_FILL
    c.border = BOX
    if fmt:
        c.number_format = fmt


def calc(c, value, fmt=N1, bold=False):
    c.value = value
    c.font = f(bold=bold)
    c.number_format = fmt


def unit(c, text):
    c.value = text
    c.font = f("595959", size=9)


def build(src=SRC, out=OUT):
    wb = load_workbook(src)
    est, asm = wb["Estimator"], wb["Assumptions"]

    # ---------------------------------------------------------------- Assumptions additions
    head(asm, 57, "Busbar ampacity model (Busbar Sizer sheet)")
    rows = [
        (58, "Ampacity coefficient k", 5.95, "A",
         "Approximate fit to DIN 43671 bare copper ratings (single bar, still air, 35 C ambient). "
         "Calibrate to your own test data"),
        (59, "Area exponent", 0.5, "", "I = k x A^0.5 x p^0.39 x (rise / reference rise)^0.61"),
        (60, "Perimeter exponent", 0.39, "", "Thin, wide bars shed more heat per mm2 of copper"),
        (61, "Temperature rise exponent", 0.61, "", "Same exponent IEC 61439-1 uses to scale current with rise"),
        (62, "Reference temperature rise for k", 30, "C", "Rise at which k was fitted"),
        (63, "Enclosure derating factor", 1.0, "ratio",
         "Placeholder. Lower it for a bar inside a closed cover (for example 0.8)"),
    ]
    for r, label, v, u, note in rows:
        asm.cell(r, 1, label).font = f()
        inp(asm.cell(r, 2), v, N2)
        asm.cell(r, 3, u).font = f()
        asm.cell(r, 4, note).font = f("595959", size=9)

    # ---------------------------------------------------------------- Keepout Zones sheet
    kz = wb.create_sheet("Keepout Zones", 1)
    kz.column_dimensions["A"].width = 46
    kz.column_dimensions["B"].width = 16
    kz.column_dimensions["C"].width = 16
    kz.column_dimensions["D"].width = 70
    kz["A1"] = "Keepout zones"
    kz["A1"].font = f(bold=True, size=14)
    kz["A2"] = ("Rectangular spaces for a bolted relay or other box. An included zone sits at the far end "
                "of the fuse row (opposite the B+ stud) and adds space to the base. Blue on yellow = input.")
    kz["A2"].font = f("595959", size=9)

    head(kz, 4, "ZONE INPUTS", 4)
    sub(kz, 5, ["", "Zone 1", "Zone 2", "Note"])
    zin = [
        (6, "Name", "Bolted relay", "Spare box", None, "Label only"),
        (7, "Include in block", "Yes", "No", None, "Yes or No"),
        (8, "Length, front to back", 60, 40, N1, "mm. Base length grows if this plus two walls is longer"),
        (9, "Width, along the fuse row", 45, 30, N1, "mm. Adds this plus one divider to the base width"),
        (10, "Height above the floor", 40, 25, N1, "mm. Wall height grows if this is taller than the fuse stack"),
        (11, "Studs", 2, 0, N0, "0, 1 or 2. Use 0 when the zone is laser welded to the bus"),
        (12, "Stud size", "M6", "M6", None, "M8, M6 or M5"),
        (13, "Current from the input bus", 100, 0, N0,
         "A. 0 = no bus connection. This current is already part of the B+ load, it is not added to it"),
        (14, "Branch length override", None, None, N1, "mm. Leave blank to use the auto length below"),
    ]
    for r, label, v1, v2, fmt, note in zin:
        kz.cell(r, 1, label).font = f()
        inp(kz.cell(r, 2), v1, fmt)
        inp(kz.cell(r, 3), v2, fmt)
        unit(kz.cell(r, 4), note)
    kz.add_data_validation(dv := DataValidation(type="list", formula1='"Yes,No"', showErrorMessage=True))
    dv.add("B7:C7")
    kz.add_data_validation(dv := DataValidation(type="list", formula1='"0,1,2"', showErrorMessage=True))
    dv.add("B11:C11")
    kz.add_data_validation(dv := DataValidation(type="list", formula1='"M8,M6,M5"', showErrorMessage=True))
    dv.add("B12:C12")
    kz.add_data_validation(dv := DataValidation(type="decimal", operator="greaterThanOrEqual",
                                                formula1="0", showErrorMessage=True))
    dv.add("B8:C10")
    dv.add("B13:C14")

    def bolt(col, field):  # Assumptions bolt table lookup on this zone's stud size
        return f"INDEX({A}${field}$31:${field}$33,MATCH({col}12,{A}$A$31:$A$33,0))"

    last_slot = "IF(Estimator!$B$20>0,Estimator!$H$20,IF(Estimator!$B$19>0,Estimator!$H$19,Estimator!$H$18))"
    last_minw = "IF(Estimator!$B$20>0,Estimator!$I$20,IF(Estimator!$B$19>0,Estimator!$I$19,Estimator!$I$18))"

    head(kz, 16, "ZONE CALCULATION", 4)
    zc = [
        (17, "Included (1 = yes)", lambda c: f'=IF({c}7="Yes",1,0)', N0, ""),
        (18, "Width added to the base", lambda c: f"={c}17*({c}9+{A}$B$6)", N1, "mm, zone width plus one divider"),
        (19, "Base length needed", lambda c: f"={c}17*({c}8+2*{A}$B$5)", N1, "mm, zone length plus two walls"),
        (20, "Wall height needed", lambda c: f"={c}17*{c}10", N1, "mm"),
        (21, "Studs counted", lambda c: f"={c}17*{c}11", N0, ""),
        (22, "Connected to the bus (1 = yes)", lambda c: f"=IF(AND({c}17=1,{c}13>0),1,0)", N0, ""),
        (23, "Branch width", lambda c: f"={c}22*MAX({c}13/({A}$B$14*Estimator!$B$7),"
                                       f"IF({c}21>0,{A}$B$16*{bolt(c, 'B')},0))", N1,
         "mm. Larger of current need (design current density) and bolt edge distance"),
        (24, "Auto branch length", lambda c: (f"={c}22*({last_slot}/2-{last_minw}/2"
                                              + ("" if c == "B" else "+$B$18") + f"+{c}18/2)"), N1,
         "mm, from the end of the input bus to the zone centre, along the row"),
        (25, "Branch length used", lambda c: f'={c}22*IF({c}14="",{c}24,{c}14)', N1, "mm"),
        (26, "Branch copper mass", lambda c: f"={c}25*{c}23*Estimator!$B$7*{A}$B$13/1000", N1,
         "g, same thickness as the bus"),
        (27, "Stud bolt and nut hardware", lambda c: f"={c}21*{bolt(c, 'I')}", MONEY, "$"),
        (28, "Stud boss size units", lambda c: f"={c}21*{bolt(c, 'J')}", N2,
         "Boss factor x studs, scaled by the retention method on the Estimator"),
        (29, "Laser weld joints (0 studs, bus connected)", lambda c: f"=IF(AND({c}22=1,{c}11=0),1,0)", N0, ""),
    ]
    for r, label, fn, fmt, note in zc:
        kz.cell(r, 1, label).font = f()
        calc(kz.cell(r, 2), fn("B"), fmt)
        calc(kz.cell(r, 3), fn("C"), fmt)
        unit(kz.cell(r, 4), note)

    head(kz, 31, "TOTALS USED BY THE ESTIMATOR", 4)
    zt = [
        (32, "Zones included", "=B17+C17", N0, ""),
        (33, "Width added to the base", "=B18+C18", N1, "mm"),
        (34, "Base length needed", "=MAX(B19,C19)", N1, "mm"),
        (35, "Wall height needed", "=MAX(B20,C20)", N1, "mm"),
        (36, "Studs retained in the base", "=B21+C21", N0, ""),
        (37, "Branch copper mass", "=B26+C26", N1, "g"),
        (38, "Stud bolt and nut hardware", "=B27+C27", MONEY, "$"),
        (39, "Stud boss size units", "=B28+C28", N2, ""),
        (40, "Laser weld joint cost", f"=(B29+C29)*{A}$B$45", MONEY, "$, Laser weld cost per joint"),
    ]
    for r, label, fm, fmt, note in zt:
        kz.cell(r, 1, label).font = f()
        calc(kz.cell(r, 2), fm, fmt, bold=True)
        unit(kz.cell(r, 4), note)

    # ---------------------------------------------------------------- Estimator edits
    est["A3"] = "Revision v1.4: keepout zones (Keepout Zones sheet) and Busbar Sizer sheet added."
    est["A3"].font = f("595959", size=9, italic=True)

    fuse_wall = "SUMPRODUCT(MAX($K$18:$K$20*$L$18:$L$20))"
    est["B36"] = f"=B34+B35+{KZ}!B33+2*{A}$B$5"
    est["B37"] = f"=MAX({fuse_wall},{KZ}!B35)"
    est["B38"] = ("=MAX(MAX(SUMPRODUCT(MAX($K$18:$K$20*$P$18:$P$20)),(B28-B27/2)+B24+Assumptions!$B$5)"
                  "+SUMPRODUCT(MAX($K$18:$K$20*$M$18:$M$20)),"
                  "INDEX(Assumptions!$F$31:$F$33,MATCH($B$10,Assumptions!$A$31:$A$33,0))+2*Assumptions!$B$5,"
                  f"{KZ}!B34)")
    est["B39"] = f"=MIN({fuse_wall},MAX(Assumptions!$B$7*{fuse_wall},B25))"
    est["A39"] = "Divider height (meets creepage, fuse stack only)"
    est["B41"] = f"=B33+B11-1+{KZ}!B32"
    est["B44"] = (f"=B36-{KZ}!B33-2*Assumptions!$B$5-IF(B20>0,H20,IF(B19>0,H19,H18))/2"
                  "+IF(B20>0,I20,IF(B19>0,I19,I18))/2")

    def style_from(dst, src):
        dst.font = copy.copy(src.font)
        dst.fill = copy.copy(src.fill)
        dst.border = copy.copy(src.border)
        dst.number_format = src.number_format
        dst.alignment = copy.copy(src.alignment)

    lbl, val, val_bold, unit_src = (copy.copy(est["A45"]), copy.copy(est["B45"]),
                                    copy.copy(est["B48"]), copy.copy(est["C45"]))
    tot_lbl, tot_val, money = copy.copy(est["A83"]), copy.copy(est["B83"]), copy.copy(est["B70"])
    head_cells = [copy.copy(est.cell(69, c)) for c in range(1, 4)]
    plain = copy.copy(est["E23"])

    def put(r, label, formula, src_val, u=None):
        est.cell(r, 1, label)
        style_from(est.cell(r, 1), lbl)
        est.cell(r, 2, formula)
        style_from(est.cell(r, 2), src_val)
        if u is not None:
            est.cell(r, 3, u)
            style_from(est.cell(r, 3), unit_src)

    # Copper block (rows 47-49 rewritten, row 49 was blank)
    put(47, "Keepout zone branch copper", f"={KZ}!B37", val, "g")
    put(48, "Copper total", "=B45+B46+B47", val, "g")
    put(49, "Copper cost", "=B48/1000*B9*(1+Assumptions!$B$15)", val_bold)
    est["C49"] = None
    style_from(est["C49"], plain)

    est["B51"] = f"=2*B33+B11+{KZ}!B36"
    est["B52"] = ("=INDEX(Assumptions!$B$37:$B$39,MATCH(B13,Assumptions!$A$37:$A$39,0))*"
                  "(2*SUMPRODUCT(B18:B20,N18:N20)+B11*INDEX(Assumptions!$J$31:$J$33,"
                  f"MATCH($B$10,Assumptions!$A$31:$A$33,0))+{KZ}!B39)")

    # Hardware and part cost block rewritten from row 69 down
    for r in range(69, 95):
        for c in range(1, 4):
            est.cell(r, c).value = None
            style_from(est.cell(r, c), plain)
    for c in range(1, 4):
        style_from(est.cell(69, c), head_cells[c - 1])
    est["A69"] = "HARDWARE AND PROCESS ADDERS"
    put(70, "Fuse bolt and nut hardware", "=2*SUMPRODUCT(B18:B20,O18:O20)", money)
    put(71, "Retention method adder (fuse, B+ and zone studs)",
        "=B51*INDEX(Assumptions!$C$37:$C$39,MATCH(B13,Assumptions!$A$37:$A$39,0))", money)
    put(72, "Joining method adder (fuse, B+ and zone studs)",
        "=B51*INDEX(Assumptions!$B$43:$B$45,MATCH(B14,Assumptions!$A$43:$A$45,0))", money)
    put(73, "B+ connection part (battery terminal or M8 stud)",
        '=IF(B12="Battery terminal",B11*Assumptions!$B$54,B11*Assumptions!$B$55)', money)
    put(74, "Keepout zone stud hardware", f"={KZ}!B38", money)
    put(75, "Keepout zone laser weld joints", f"={KZ}!B40", money)

    for c in range(1, 4):
        style_from(est.cell(77, c), head_cells[c - 1])
    est["A77"] = "PART COST"
    cost_rows = [("Plastic base", "=B59"), ("Cover", "=B67"), ("Copper", "=B49"),
                 ("Fuse bolt and nut hardware", "=B70"), ("B+ connection part", "=B73"),
                 ("Retention adder", "=B71"), ("Joining adder", "=B72"),
                 ("Keepout zone hardware and welds", "=B74+B75")]
    for i, (label, fm) in enumerate(cost_rows):
        put(78 + i, label, fm, money)
    est["A86"], est["B86"] = "Total part cost", "=SUM(B78:B85)"
    style_from(est["A86"], tot_lbl)
    style_from(est["B86"], tot_val)
    put(87, "Cost per fuse position", "=B86/B33", money)
    put(88, "Total plastic mass (base and cover)", "=B58+B66", val, "g")

    est["E5"], est["E7"], est["E8"] = "=B86", "=B88", "=B48"
    est["D12"], est["E12"] = "Keepout zones included", f"={KZ}!B32"
    style_from(est["D12"], est["D11"])
    style_from(est["E12"], est["E11"])
    est["E12"].number_format = N0

    # openpyxl drops the x14 (cross-sheet) list validations on load, so add them back
    for ref, rng in (("B13", "Assumptions!$A$37:$A$39"), ("B14", "Assumptions!$A$43:$A$45")):
        dv = DataValidation(type="list", formula1=rng, showErrorMessage=True)
        dv.add(ref)
        est.add_data_validation(dv)

    # ---------------------------------------------------------------- Busbar Sizer sheet
    bs = wb.create_sheet("Busbar Sizer", 2)
    for col, w in zip("ABCDEFGH", (40, 14, 14, 14, 14, 14, 14, 14)):
        bs.column_dimensions[col].width = w
    bs["A1"] = "Busbar Sizer"
    bs["A1"].font = f(bold=True, size=14)
    bs["A2"] = ("Quick copper cost for a known bar, and the cheapest width-to-thickness ratio at a fixed "
                "current. Ampacity model and its constants are on the Assumptions sheet (rows 57 to 63).")
    bs["A2"].font = f("595959", size=9)

    head(bs, 4, "INPUTS")
    bin_ = [(5, "Current", 300, N0, "A"), (6, "Allowed temperature rise", 30, N1, "C above ambient"),
            (7, "Bar length", 150, N1, "mm"), (8, "Known bar width", 40, N1, "mm"),
            (9, "Known bar thickness", 2, N2, "mm"), (11, "Max width that fits the package", 30, N1, "mm")]
    for r, label, v, fmt, u in bin_:
        bs.cell(r, 1, label).font = f()
        inp(bs.cell(r, 2), v, fmt)
        unit(bs.cell(r, 3), u)
    bs["A10"] = "Copper price"
    bs["A10"].font = f()
    calc(bs["B10"], "=Estimator!B9", MONEY)
    bs["B10"].font = f("008000")
    unit(bs["C10"], "$/kg, linked to the Estimator")

    k = f"{A}$B$58*{A}$B$63"
    ea, ep, en, eref = f"{A}$B$59", f"{A}$B$60", f"{A}$B$61", f"{A}$B$62"

    head(bs, 12, "KNOWN BAR")
    kb = [
        (13, "Width to thickness ratio", "=B8/B9", N1, ": 1"),
        (14, "Cross-section area", "=B8*B9", N1, "mm2"),
        (15, "Perimeter", "=2*(B8+B9)", N1, "mm"),
        (16, "Current density", "=B5/B14", N2, "A/mm2"),
        (17, "Rated current at the allowed rise", f"={k}*B14^{ea}*B15^{ep}*(B6/{eref})^{en}", N1, "A"),
        (18, "Estimated temperature rise at the current", f"={eref}*(B5/({k}*B14^{ea}*B15^{ep}))^(1/{en})",
         N1, "C"),
        (19, "Carries the current?", '=IF(B17>=B5,"Yes","No")', None, ""),
        (20, "Mass", f"=B7*B14*{A}$B$13/1000", N1, "g"),
        (21, "Copper cost (with plating uplift)", f"=B20/1000*B10*(1+{A}$B$15)", MONEY, "$"),
    ]
    for r, label, fm, fmt, u in kb:
        bs.cell(r, 1, label).font = f()
        calc(bs.cell(r, 2), fm, fmt or "General", bold=(r == 21))
        unit(bs.cell(r, 3), u)

    head(bs, 23, "RATIO SWEEP AT FIXED CURRENT", 8)
    bs["A24"] = ("Each row is the smallest bar that carries the current at the allowed rise for that ratio. "
                 "Edit the ratios in column A.")
    bs["A24"].font = f("595959", size=9)
    sub(bs, 25, ["Width : thickness", "Thickness (mm)", "Width (mm)", "Area (mm2)", "A/mm2",
                 "Mass (g)", "Cost ($)", "Fits? (0.8 to 5.0 mm thick, width within max)"])
    ratios = [1, 1.5, 2, 3, 4, 5, 6, 8, 10, 12, 15, 20]
    first, last = 26, 26 + len(ratios) - 1
    ieff = f"($B$5/({k}*($B$6/{eref})^{en}))"
    for i, rr in enumerate(ratios):
        r = first + i
        inp(bs.cell(r, 1), rr, N1)
        calc(bs.cell(r, 2), f"=({ieff}/(A{r}^{ea}*(2*(A{r}+1))^{ep}))^(1/(2*{ea}+{ep}))", N2)
        calc(bs.cell(r, 3), f"=A{r}*B{r}", N1)
        calc(bs.cell(r, 4), f"=B{r}*C{r}", N1)
        calc(bs.cell(r, 5), f"=$B$5/D{r}", N2)
        calc(bs.cell(r, 6), f"=$B$7*D{r}*{A}$B$13/1000", N1)
        calc(bs.cell(r, 7), f"=F{r}/1000*$B$10*(1+{A}$B$15)", MONEY)
        calc(bs.cell(r, 8), f'=IF(AND(B{r}>=0.8,B{r}<=5,C{r}<=$B$11),"Yes","No")', "General")
        for c in range(2, 9):
            bs.cell(r, c).border = BOX
    r = last + 2
    bs.cell(r, 1, "Cheapest ratio that fits").font = f(bold=True)
    calc(bs.cell(r, 2), f'=IFERROR(INDEX(A{first}:A{last},MATCH(_xlfn.MINIFS(G{first}:G{last},H{first}:H{last},"Yes"),'
                        f'G{first}:G{last},0)),"None fit")', N1, bold=True)
    bs.cell(r, 2).fill = RESULT_FILL
    unit(bs.cell(r, 3), ": 1")
    bs.cell(r + 1, 1, "Its cost").font = f()
    calc(bs.cell(r + 1, 2), f'=IF(COUNTIF(H{first}:H{last},"Yes")=0,"-",_xlfn.MINIFS(G{first}:G{last},H{first}:H{last},"Yes"))', MONEY)
    bs.cell(r + 2, 1, ("Cost keeps falling as the bar gets thinner and wider, so the cheapest bar is the widest "
                       "one that fits. Stiffness and the 0.8 mm minimum are the other limits.")).font = f("595959", size=9)

    ch = ScatterChart()
    ch.title = "Copper cost vs width : thickness"
    ch.style = 13
    ch.x_axis.title = "Width : thickness"
    ch.y_axis.title = "Cost ($)"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    s = Series(Reference(bs, min_col=7, min_row=first, max_row=last),
               Reference(bs, min_col=1, min_row=first, max_row=last), title="Cost ($)")
    s.marker.symbol = "circle"
    ch.series.append(s)
    ch.height, ch.width = 8, 16
    bs.add_chart(ch, "J4")

    wb.save(out)
    return out


if __name__ == "__main__":
    print(build(*sys.argv[1:]))
