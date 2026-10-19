"""Builds Battery_Fuse_Block_Cost_Estimator_v1.9.xlsx on top of the v1.8 build.

v1.9 limits busbar thickness to copper mill stock sizes (Mike, 2026-10-19):
  * Stock thickness list on Assumptions A83:A94 (placeholder 0.8 to 5.0 mm, 12 slots, blank = unused).
  * Estimator B7 (copper thickness) and Busbar Sizer B9 (known bar thickness) pick from that list.
  * The Busbar Sizer sweep runs over the stock thicknesses instead of width : thickness ratios.
    Each row is the narrowest bar of that thickness that carries the current at the allowed rise
    (6 Newton steps in columns L to R), and the ratio is shown as a result. Fits = width within max.

The web tool (bft/web/bft_estimator_v1.9.html) must match these formulas.
Run: python bft/tools/build_v1_9.py
"""
import os
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_v1_7 as v7  # noqa: E402
import build_v1_8 as v8  # noqa: E402
from build_v1_4 import N1, N2, calc, f, head, inp, sub, unit  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "Battery_Fuse_Block_Cost_Estimator_v1.9.xlsx")
A = "Assumptions!"
STOCK = [0.8, 1.0, 1.2, 1.5, 2.0, 2.5, 3.0, 4.0, 5.0, None, None, None]
FIRST, LAST = 26, 37
STEPS = "MNOPQR"  # Newton steps; L is the start value


def sweep_formulas(r):
    blank = f'$A{r}=""'
    t, a, p = f"$A{r}", f"{A}$B$59", f"{A}$B$60"
    out = {"A": f'=IF({A}$A${83 + r - FIRST}="","",{A}$A${83 + r - FIRST})',
           "L": f'=IF({blank},"",$B$5/({A}$B$74*{t}))'}
    prev = "L"
    for col in STEPS:
        w = f"{prev}{r}"
        out[col] = (f'=IF({blank},"",IF($B$5<=0,0,{w}*EXP(-({a}*LN({w}*{t})+{p}*LN(2*({w}+{t}))-LN($K$25))'
                    f"/({a}+{p}*{w}/({w}+{t})))))")
        prev = col
    out.update({
        "B": f'=IF({blank},"",R{r})',
        "C": f'=IF({blank},"",B{r}/A{r})',
        "D": f'=IF({blank},"",A{r}*B{r})',
        "E": f'=IF({blank},"",$B$5/D{r})',
        "F": f'=IF({blank},"",$B$7*D{r}*{A}$B$13/1000)',
        "G": f'=IF({blank},"",F{r}/1000*$B$10*(1+{A}$B$15))',
        "H": f'=IF({blank},"",IF(B{r}<=$B$11,"Yes","No"))',
        "I": f'=IF(H{r}="Yes",G{r},"")',
    })
    return out


def edit(wb):
    est, asm, bs = wb["Estimator"], wb["Assumptions"], wb["Busbar Sizer"]

    # ---------------------------------------------------------------- Assumptions: stock list
    head(asm, 82, "Copper mill stock thicknesses (mm)")
    for i, v in enumerate(STOCK):
        c = asm.cell(83 + i, 1)
        inp(c, v, "0.00")
    unit(asm["B83"], "Placeholder list. Edit to match your copper supplier. Blank = unused. "
                     "Estimator B7 and Busbar Sizer B9 pick from it, and the Busbar Sizer sweeps it.")

    # ---------------------------------------------------------------- thickness inputs pick from the list
    est.data_validations.dataValidation = [d for d in est.data_validations.dataValidation
                                           if "B7" not in str(d.sqref).split()]
    for ws, cell in ((est, "B7"), (bs, "B9")):
        d = DataValidation(type="list", formula1=f"{A}$A$83:$A$94", showErrorMessage=True,
                           error="Pick a mill stock thickness (Assumptions A83:A94).")
        d.add(cell)
        ws.add_data_validation(d)
    est["C7"] = "mm, mill stock (Assumptions A83:A94)"
    bs["C9"] = "mm, mill stock (Assumptions A83:A94)"
    est["A3"] = ("Revision v1.9: busbar thickness limited to copper mill stock sizes (Assumptions A83:A94). "
                 "Builds on v1.8.")

    # ---------------------------------------------------------------- Busbar Sizer: stock thickness sweep
    bs.cell(2, 1).value = ("Quick copper cost for a known bar, and the cheapest mill stock thickness at a fixed "
                           "current. Ampacity model: Assumptions rows 57 to 63 and 69 to 73.")
    bs["A23"] = "STOCK THICKNESS SWEEP AT FIXED CURRENT"
    bs["A24"] = ("Each row is the narrowest bar of that stock thickness that carries the current at the allowed "
                 "rise. Edit the stock list on Assumptions A83:A94.")
    sub(bs, 25, ["Stock thickness (mm)", "Width needed (mm)", "Width : thickness", "Area (mm2)", "A/mm2",
                 "Mass (g)", "Cost ($)", "Fits? (width within max)", "Cost if it fits ($)"])
    bs["K24"] = "Load over k at the allowed rise"
    bs["K24"].font = f(size=9)
    calc(bs["K25"], f"=$B$5/({v7.K_NEW}*($B$6/{A}$B$62)^{A}$B$61)", "0.000")
    bs["L24"] = "Width solve (start, then 6 Newton steps), mm"
    bs["L24"].font = f(size=9)
    fmts = {"A": "0.00", "B": N1, "C": N1, "D": N1, "E": N2, "F": N1, "G": "$0.00", "I": "$0.00"}
    for r in range(FIRST, LAST + 1):
        for col, fm in sweep_formulas(r).items():
            c = bs[f"{col}{r}"]
            c.value = fm
            c.number_format = fmts.get(col, N1 if col in "L" + STEPS else "General")
            c.font = f(size=9) if col in "L" + STEPS else f()
            c.fill = PatternFill(fill_type=None)  # column A is now calculated, not an input
    bs["A39"] = "Cheapest stock thickness that fits"
    bs["B39"] = '=IF(COUNT(I26:I37)=0,"None fit",INDEX(A26:A37,MATCH(MIN(I26:I37),I26:I37,0)))'
    bs["C39"] = "mm"
    bs["A41"] = "Its width"
    bs["B41"] = '=IF(COUNT(I26:I37)=0,"-",INDEX(B26:B37,MATCH(MIN(I26:I37),I26:I37,0)))'
    bs["B41"].number_format = N1
    bs["C41"] = "mm"
    bs["A42"] = ("Cost keeps falling as the bar gets thinner and wider, so the cheapest bar is usually the "
                 "thinnest stock that fits the max width. Stiffness is the other limit.")
    bs["A42"].font = f("595959", size=9)

    bs._charts = []
    ch = ScatterChart()
    ch.title = "Copper cost vs stock thickness"
    ch.scatterStyle = "lineMarker"
    ch.style = 2
    ch.legend = None
    ch.height, ch.width = 9, 17
    ch.x_axis.title = "Stock thickness (mm)"
    ch.y_axis.title = "Copper cost ($)"
    ch.x_axis.axPos, ch.y_axis.axPos = "b", "l"
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.y_axis.number_format = "$0.00"
    ch.x_axis.number_format = "0.0"
    ch.x_axis.scaling.min = ch.y_axis.scaling.min = 0
    s = Series(Reference(bs, min_col=7, min_row=FIRST, max_row=LAST),
               Reference(bs, min_col=1, min_row=FIRST, max_row=LAST), title="Cost ($)")
    s.graphicalProperties.line.solidFill = "2F6FA8"
    s.graphicalProperties.line.width = 28575
    s.marker.symbol = "circle"
    s.marker.size = 7
    s.marker.graphicalProperties.solidFill = "2F6FA8"
    s.marker.graphicalProperties.line.solidFill = "2F6FA8"
    s.smooth = False
    ch.series.append(s)
    bs.add_chart(ch, "J4")


def build(out=OUT, cache=True):
    tmp = os.path.join(tempfile.mkdtemp(), "v18.xlsx")
    v8.build(tmp, cache=False)
    wb = load_workbook(tmp)
    edit(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    if cache:
        v7.add_cached_values(out)
    return out


if __name__ == "__main__":
    print(build(*sys.argv[1:]))
