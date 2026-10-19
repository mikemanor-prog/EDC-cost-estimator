"""Builds Battery_Fuse_Block_Cost_Estimator_v1.5.xlsx on top of the v1.4 build.

v1.5 replaces the Keepout Zones sheet with a Layout sheet:
  * Keepout zones are sealed boxes. Base walls stay at fuse height and the cover bumps out
    over a zone that is taller than the walls. Each zone sits at the end of the row or on
    the back side (behind the bus).
  * Fuses can be moved to the short side. They turn 90 degrees and the bus becomes an L.
    The exit end (opposite B+ or the B+ end) is an input.
  * Fuse slot pitch can be overridden per fuse type, and the walls between fuses can be
    switched off (pitch then needs body width + creepage instead of a divider).
  * Max block length and depth on the battery are checked (flag only).

The web tool (bft/web/bft_estimator_v1.5.html) must match these formulas.
Run: python bft/tools/build_v1_5.py
"""
import os
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_v1_4 as v4  # noqa: E402
from build_v1_4 import MONEY, N0, N1, N2, calc, f, head, inp, sub, unit  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "Battery_Fuse_Block_Cost_Estimator_v1.5.xlsx")

L = "Layout!"
A = "Assumptions!"
E = "Estimator!"
WALL, DIV, FIT, COVT = f"{A}$B$5", f"{A}$B$6", f"{A}$B$48", f"{A}$B$8"


def dv(ws, kind, rng, formula1=None, **kw):
    d = DataValidation(type=kind, formula1=formula1, showErrorMessage=True, **kw)
    d.add(rng)
    ws.add_data_validation(d)


def build(out=OUT):
    tmp = os.path.join(tempfile.mkdtemp(), "v14.xlsx")
    v4.build(v4.SRC, tmp)
    wb = load_workbook(tmp)
    del wb["Keepout Zones"]
    est = wb["Estimator"]

    ly = wb.create_sheet("Layout", 1)
    for col, w in zip("ABCDEFGHI", (44, 14, 14, 14, 14, 14, 14, 30, 12)):
        ly.column_dimensions[col].width = w
    ly["A1"] = "Layout"
    ly["A1"].font = f(bold=True, size=14)
    ly["A2"] = ("Fuse placement, slot pitch, keepout zones and battery fit. Blue on yellow = input. "
                "Long side = the side along the fuse row. Short side = the end opposite B+ or the B+ end.")
    ly["A2"].font = f("595959", size=9)

    # ------------------------------------------------------------------ layout inputs
    head(ly, 4, "LAYOUT INPUTS", 4)
    rows = [(5, "Walls between fuses", "Yes", None, "Yes or No. No = sockets may overlap, pitch needs body + creepage"),
            (6, "Short-side outputs exit at", "Opposite B+", None, "Opposite B+ or B+ end"),
            (7, "Max block length along the row", None, N1, "mm, cover outside. Blank = no limit"),
            (8, "Max block depth", None, N1, "mm, cover outside. Blank = no limit")]
    for r, label, v, fmt, note in rows:
        ly.cell(r, 1, label).font = f()
        inp(ly.cell(r, 2), v, fmt)
        unit(ly.cell(r, 3), note)
    dv(ly, "list", "B5", '"Yes,No"')
    dv(ly, "list", "B6", '"Opposite B+,B+ end"')

    head(ly, 10, "FUSE PLACEMENT AND PITCH", 9)
    sub(ly, 11, ["Fuse type", "Qty", "On short side", "Pitch override (mm)", "Auto pitch (mm)",
                 "Pitch used (mm)", "On long side", "Pitch check", "Short used"])
    sock = "INDEX(Assumptions!$F$31:$F$33,MATCH(Estimator!C{r},Assumptions!$A$31:$A$33,0))"
    for i in range(3):
        r, er = 12 + i, 18 + i
        calc(ly.cell(r, 1), f"={E}A{er}", "General")
        calc(ly.cell(r, 2), f"={E}B{er}", N0)
        inp(ly.cell(r, 3), 0, N0)
        inp(ly.cell(r, 4), None, N1)
        calc(ly.cell(r, 5), (f'=IF($B$5="Yes",MAX({E}E{er},{sock.format(r=er)},{E}I{er}+{E}$B$24)+{DIV},'
                             f"MAX({E}E{er}+{E}$B$25,{E}I{er}+{E}$B$24))"), N1)
        calc(ly.cell(r, 6), f'=IF(D{r}="",E{r},D{r})', N1)
        calc(ly.cell(r, 7), f"=B{r}-I{r}", N0)
        calc(ly.cell(r, 8), f'=IF(F{r}<E{r},"Below auto: check socket access and creepage","OK")', "General")
        calc(ly.cell(r, 9), f"=MIN(B{r},C{r})", N0)
    dv(ly, "whole", "C12:C14", "0", operator="greaterThanOrEqual")
    dv(ly, "decimal", "D12:D14", "0", operator="greaterThan")
    ly["A15"] = ("Auto pitch with walls = largest of body width, socket clearance and bus edge width + clearance, "
                 "plus one divider. Without walls = largest of body width + creepage and bus edge width + clearance.")
    ly["A15"].font = f("595959", size=9)

    # ------------------------------------------------------------------ zone inputs
    head(ly, 17, "KEEPOUT ZONE INPUTS (SEALED BOXES)", 4)
    sub(ly, 18, ["", "Zone 1", "Zone 2", "Note"])
    zin = [(19, "Name", "Bolted relay", "Spare box", None, "Label only"),
           (20, "Include in block", "Yes", "No", None, "Yes or No"),
           (21, "Placement", "End of row", "Back side", None, "End of row, or Back side (behind the bus)"),
           (22, "Length, front to back", 60, 40, N1, "mm"),
           (23, "Width, along the fuse row", 45, 30, N1, "mm"),
           (24, "Height above the floor", 40, 25, N1, "mm. Cover bumps out if taller than the base wall"),
           (25, "Studs", 2, 0, N0, "0, 1 or 2. Use 0 when the zone is laser welded to the bus"),
           (26, "Stud size", "M6", "M6", None, "M8, M6 or M5"),
           (27, "Current from the input bus", 100, 0, N0,
            "A. 0 = no bus connection. Already part of the B+ load, not added to it"),
           (28, "Branch length override", None, None, N1, "mm. Blank = auto length below")]
    for r, label, v1, v2, fmt, note in zin:
        ly.cell(r, 1, label).font = f()
        inp(ly.cell(r, 2), v1, fmt)
        inp(ly.cell(r, 3), v2, fmt)
        unit(ly.cell(r, 4), note)
    dv(ly, "list", "B20:C20", '"Yes,No"')
    dv(ly, "list", "B21:C21", '"End of row,Back side"')
    dv(ly, "list", "B25:C25", '"0,1,2"')
    dv(ly, "list", "B26:C26", '"M8,M6,M5"')

    def bolt(c, field):
        return f"INDEX({A}${field}$31:${field}$33,MATCH({c}26,{A}$A$31:$A$33,0))"

    head(ly, 30, "KEEPOUT ZONE CALCULATION", 4)
    zc = [
        (31, "Included (1 = yes)", lambda c: f'=IF({c}20="Yes",1,0)', N0, ""),
        (32, "At the end of the row (1 = yes)", lambda c: f'={c}31*IF({c}21="End of row",1,0)', N0, ""),
        (33, "On the back side (1 = yes)", lambda c: f"={c}31-{c}32", N0, ""),
        (34, "Width added to the row", lambda c: f"={c}32*({c}23+{DIV})", N1, "mm, zone width plus one divider"),
        (35, "Width needed on the back side", lambda c: f"={c}33*({c}23+{DIV})", N1, "mm"),
        (36, "Depth added on the back side", lambda c: f"={c}33*({c}22+{DIV})", N1,
         "mm, zone length plus one divider in front of it"),
        (37, "Base length needed (end of row)", lambda c: f"={c}32*({c}22+2*{WALL}+IF($B$55=1,$B$62-{WALL},0))",
         N1, "mm. Starts in front of the bus when the bus runs past it to the short side"),
        (38, "Cover bump-out height", lambda c: f"={c}31*MAX(0,{c}24+{FIT}-{E}$B$37)", N1,
         "mm above the base wall"),
        (39, "Cover bump-out volume", lambda c: f"=2*({c}22+{c}23)*{COVT}*{c}38", N0, "mm3, side walls of the bump"),
        (40, "Studs counted", lambda c: f"={c}31*{c}25", N0, ""),
        (41, "Connected to the bus (1 = yes)", lambda c: f"=IF(AND({c}31=1,{c}27>0),1,0)", N0, ""),
        (42, "Branch width", lambda c: (f"={c}41*MAX({c}27/({A}$B$14*{E}$B$7),"
                                         f"IF({c}40>0,{A}$B$16*{bolt(c, 'B')},0))"), N1,
         "mm. Larger of current need and bolt edge distance"),
        (43, "Auto branch length", lambda c: (f"={c}41*IF({c}33=1,{c}22/2+{DIV}+{E}$B$24,IF($B$55=1,{c}22/2,"
                                               f"$B$56/2-$B$57/2" + ("" if c == "B" else "+$B$34")
                                               + f"+{c}34/2))"), N1,
         "mm. Back side: bus to zone centre. End of row: bus end to zone centre"),
        (44, "Branch length used", lambda c: f'={c}41*IF({c}28="",{c}43,{c}28)', N1, "mm"),
        (45, "Branch copper mass", lambda c: f"={c}44*{c}42*{E}$B$7*{A}$B$13/1000", N1, "g, bus thickness"),
        (46, "Stud bolt and nut hardware", lambda c: f"={c}40*{bolt(c, 'I')}", MONEY, "$"),
        (47, "Stud boss size units", lambda c: f"={c}40*{bolt(c, 'J')}", N2, ""),
        (48, "Laser weld joints (0 studs, bus connected)", lambda c: f"=IF(AND({c}41=1,{c}25=0),1,0)", N0, ""),
    ]
    for r, label, fn, fmt, note in zc:
        ly.cell(r, 1, label).font = f()
        calc(ly.cell(r, 2), fn("B"), fmt)
        calc(ly.cell(r, 3), fn("C"), fmt)
        unit(ly.cell(r, 4), note)

    # ------------------------------------------------------------------ short side and bus
    head(ly, 50, "SHORT SIDE AND BUS", 4)
    kl = "(G12:G14>0)"
    ks = "(I12:I14>0)"
    sb = [
        (51, "Fuses on the short side", "=SUM(I12:I14)", N0, ""),
        (52, "Fuses on the long side", "=SUM(G12:G14)", N0, ""),
        (53, "Short-side slots length", "=SUMPRODUCT(I12:I14,F12:F14)", N1, "mm"),
        (54, "Long-side slots width", "=SUMPRODUCT(G12:G14,F12:F14)", N1, "mm"),
        (55, "Bus runs past the end zones to the short side (1 = yes)",
         '=IF(AND(B51>0,B6="Opposite B+"),1,0)', N0, ""),
        (56, "Last long-side fuse pitch", "=IF(G14>0,F14,IF(G13>0,F13,F12))", N1, "mm"),
        (57, "Last long-side fuse min bus width",
         f"=IF(G14>0,{E}I20,IF(G13>0,{E}I19,{E}I18))", N1, "mm"),
        (58, "Last short-side fuse pitch", "=IF(I14>0,F14,IF(I13>0,F13,F12))", N1, "mm"),
        (59, "Last short-side fuse min bus width", f"=IF(I14>0,{E}I20,IF(I13>0,{E}I19,{E}I18))", N1, "mm"),
        (60, "Back-side zone depth", "=MAX(B36,C36)", N1, "mm"),
        (61, "Input stud row from the back edge",
         f"=B60+MAX(SUMPRODUCT(MAX({kl}*{E}P18:P20)),({E}B28-{E}B27/2)+{E}B24+{WALL})", N1, "mm"),
        (62, "Bus front edge from the back edge", f"=B61+{E}B27/2", N1, "mm"),
        (63, "Short-side input stud offset",
         f"=IF(B51>0,MAX(SUMPRODUCT(MAX({ks}*({E}P18:P20-{WALL}))),{E}B28-{E}B27/2+{E}B24),0)", N1,
         "mm, from the inner edge of the short-side column"),
        (64, "Short-side column width",
         f"=IF(B51>0,B63+SUMPRODUCT(MAX({ks}*({E}M18:M20-{WALL}))),0)", N1, "mm, outer wall not included"),
        (65, "Long-row bus end (from the B+ end inner wall)",
         f"=IF(B52>0,{E}B35+B54-B56/2+B57/2,{E}B35)", N1, "mm"),
        (66, "Bus length along the row",
         f'=IF(B51=0,B65,IF(B6="Opposite B+",{E}B35+B54+B34+C34+B63+{E}B27/2,B65+B63+{E}B27/2))', N1, "mm"),
        (67, "Bus length down the short side", "=IF(B51>0,B53-B58/2+B59/2,0)", N1, "mm, past the bus front edge"),
        (68, "Base length needed for the short side", f"=IF(B51>0,B62+B53+{WALL},0)", N1, "mm"),
        (69, "Short-side divider count", '=IF(B51>0,IF(B5="Yes",B51,1),0)', N0, ""),
        (70, "Long-side divider count", f'=IF(B5="Yes",MAX(0,B52+{E}B11-1),{E}B11)', N0, ""),
        (71, "Divider length, total",
         f"=(B70+B32+C32)*({E}B38-2*{WALL})+B69*B64+B23*B33+C23*C33", N1, "mm"),
        (72, "Divider count, total", "=B70+B69+B31+C31", N0, ""),
    ]
    for r, label, fm, fmt, note in sb:
        ly.cell(r, 1, label).font = f()
        calc(ly.cell(r, 2), fm, fmt)
        unit(ly.cell(r, 3), note)

    head(ly, 74, "TOTALS USED BY THE ESTIMATOR", 4)
    tot = [
        (75, "Zones included", "=B31+C31", N0, ""),
        (76, "Width added to the row by end zones", "=B34+C34", N1, "mm"),
        (77, "Width needed by back-side zones", "=B35+C35", N1, "mm, includes one divider per zone"),
        (78, "Base length needed by end zones", "=MAX(B37,C37)", N1, "mm"),
        (79, "Inside height needed by zones", f"=MAX(B31*(B24+{FIT}),C31*(C24+{FIT}))", N1, "mm"),
        (80, "Cover bump-out volume", "=B39+C39", N0, "mm3"),
        (81, "Zone studs retained in the base", "=B40+C40", N0, ""),
        (82, "Zone branch copper", "=B45+C45", N1, "g"),
        (83, "Zone stud bolt and nut hardware", "=B46+C46", MONEY, "$"),
        (84, "Zone stud boss size units", "=B47+C47", N2, ""),
        (85, "Laser weld joint cost", f"=(B48+C48)*{A}$B$45", MONEY, "$"),
        (86, "Battery length check",
         f'=IF(B7="","No limit",IF({E}E66<=B7,"Fits","Too long by "&TEXT({E}E66-B7,"0.0")&" mm"))', None, "cover outside"),
        (87, "Battery depth check",
         f'=IF(B8="","No limit",IF({E}E65<=B8,"Fits","Too deep by "&TEXT({E}E65-B8,"0.0")&" mm"))', None, "cover outside"),
    ]
    for r, label, fm, fmt, note in tot:
        ly.cell(r, 1, label).font = f()
        calc(ly.cell(r, 2), fm, fmt or "General", bold=True)
        unit(ly.cell(r, 3), note)

    # ------------------------------------------------------------------ Estimator rewiring
    est["A3"] = ("Revision v1.5: Layout sheet (sealed keepout zones with cover bump-out, short-side outputs, "
                 "pitch override, walls on/off, battery fit). Busbar Sizer sheet from v1.4.")
    est["H17"] = "Slot pitch"
    for i in range(3):
        est[f"H{18 + i}"] = f"={L}F{12 + i}"
    est["A34"] = "Long-side fuse slots width"
    est["B34"] = f"={L}B54"
    est["B36"] = f"=MAX(B34+B35+{L}B76+{L}B64+2*{WALL},{L}B77-{DIV}+2*{WALL})"
    est["B37"] = "=SUMPRODUCT(MAX($K$18:$K$20*$L$18:$L$20))"
    est["B38"] = (f"=MAX({L}B61+SUMPRODUCT(MAX(({L}G12:G14>0)*$M$18:$M$20)),"
                  f"INDEX(Assumptions!$F$31:$F$33,MATCH($B$10,Assumptions!$A$31:$A$33,0))+2*{WALL}+{L}B60,"
                  f"{L}B68,{L}B78)")
    est["A39"] = "Divider height (meets creepage)"
    est["B39"] = "=MIN(B37,MAX(Assumptions!$B$7*B37,B25))"
    est["A40"] = "Overall height with cover (incl. zone bump-out)"
    est["B40"] = f"=Assumptions!$B$4+MAX(B37,{L}B79)+Assumptions!$B$8"
    est["B41"] = f"={L}B72"
    est["A44"] = "Bus length (along the row plus down the short side)"
    est["B44"] = f"={L}B66+{L}B67"
    est["B47"] = f"={L}B82"
    est["B51"] = f"=2*B33+B11+{L}B81"
    est["B52"] = ("=INDEX(Assumptions!$B$37:$B$39,MATCH(B13,Assumptions!$A$37:$A$39,0))*"
                  "(2*SUMPRODUCT(B18:B20,N18:N20)+B11*INDEX(Assumptions!$J$31:$J$33,"
                  f"MATCH($B$10,Assumptions!$A$31:$A$33,0))+{L}B84)")
    est["A55"] = "Divider volume (total divider length x thickness x height)"
    est["B55"] = f"={L}B71*Assumptions!$B$6*B39"
    est["A65"] = "Cover volume with scrap (incl. zone bump-out)"
    est["B65"] = f"=(B62+B63+B64+{L}B80)*(1+Assumptions!$B$12)"
    est["E66"] = (f'=B36+2*E62+IF(AND({L}B51>0,{L}B6="B+ end"),MAX(E63,E64),E63+IF({L}B51>0,E64,0))')
    est["B74"] = f"={L}B83"
    est["B75"] = f"={L}B85"
    est["E12"] = f"={L}B75"
    for r, label, src in ((13, "Fits battery length?", "B86"), (14, "Fits battery depth?", "B87")):
        est.cell(r, 4, label)
        est.cell(r, 5, f"={L}{src}")
        v4_style(est.cell(r, 4), est["D11"])
        v4_style(est.cell(r, 5), est["E11"])

    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "Keepout Zones" in c.value:
                    raise RuntimeError(f"stale reference in {ws.title}!{c.coordinate}: {c.value}")
    wb.save(out)
    return out


def v4_style(dst, src):
    import copy
    dst.font = copy.copy(src.font)
    dst.fill = copy.copy(src.fill)
    dst.border = copy.copy(src.border)
    dst.alignment = copy.copy(src.alignment)


if __name__ == "__main__":
    print(build(*sys.argv[1:]))
