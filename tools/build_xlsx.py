"""Builds EDC_Cost_Estimator.xlsx: an Excel port of the EDC Plastic Cost Estimator web page.

All results are live formulas. Run: python tools/build_xlsx.py
"""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "EDC_Cost_Estimator.xlsx"
FONT = "Arial"
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD_FILL = PatternFill("solid", fgColor="1F2A3A")
SUB_FILL = PatternFill("solid", fgColor="DCE3EC")
TOT_FILL = PatternFill("solid", fgColor="FCE9C8")
THIN = Side(style="thin", color="B7BFCB")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

MONEY = '$#,##0.0000;($#,##0.0000);"-"'
MONEY2 = '$#,##0.00;($#,##0.00);"-"'
N0, N1, N2 = "#,##0", "#,##0.0", "#,##0.00"

PARTS = [  # id, name, kind, L, W, H, example count
    ("micro2", "Micro2", "fuse", 11.5, 6, 16, 20),
    ("micro3", "Micro3", "fuse", 17, 6.5, 15, 10),
    ("mcase", "M-Case", "fuse", 15, 10, 14, 6),
    ("a1s", "A1S (placeholder size)", "fuse", 20, 12, 18, 2),
    ("relMicro", "Micro PCB relay (up to 35 A)", "relay", 16, 14, 15, 8),
    ("relMini", "Mini PCB relay (35-70 A)", "relay", 24, 18, 22, 4),
    ("relHigh", "High-current PCB relay (70 A+)", "relay", 32, 24, 28, 0),
]
RESINS = [("PA66-GF30", 1.36, 3.8), ("PBT-GF30", 1.53, 3.2), ("PP-T20", 1.05, 1.8)]
GEO = [
    ("pack", "Packing efficiency (0-1)", 0.65, "0.00"),
    ("conn", "Connector / stud area (mm²)", 2500, N0),
    ("edge", "Edge margin (mm)", 6, N1),
    ("aspect", "Aspect ratio L:W", 1.5, N2),
    ("wall", "Housing wall (mm)", 2.5, N1),
    ("rib", "Cavity rib (mm)", 1.2, N1),
    ("stack", "PCB + busbar stack (mm)", 10, N1),
    ("coverWall", "Cover wall (mm)", 2, N1),
    ("fuseHead", "Fuse head above housing (mm)", 8, N1),
    ("coverClear", "Cover clearance (mm)", 6, N1),
    ("minLower", "Min lower cover height (mm)", 15, N1),
]
MOLD = [
    ("scrap", "Runner / scrap (%)", 0.08, "0.0%"),
    ("base", "Cycle base (s)", 12, N1),
    ("k", "Cooling k (s/mm²)", 2, N1),
    ("clamp", "Clamp (tons/in²)", 3.5, N1),
    ("r300", "Press rate <300 t ($/hr)", 45, MONEY2),
    ("r500", "Press rate 300-500 t ($/hr)", 60, MONEY2),
    ("r800", "Press rate 500-800 t ($/hr)", 80, MONEY2),
    ("rMax", "Press rate >800 t ($/hr)", 110, MONEY2),
]


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def header(ws, row, text, span):
    ws.cell(row, 1, text).font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
    for c in range(1, span + 1):
        ws.cell(row, c).fill = HEAD_FILL


def colheads(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font = font(bold=True)
        c.fill = SUB_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center" if i else "left", wrap_text=True)


def inp(cell, value, fmt=None):
    cell.value = value
    cell.font = font(BLUE)
    cell.fill = YELLOW
    cell.border = BOX
    if fmt:
        cell.number_format = fmt


wb = Workbook()
est = wb.active
est.title = "Estimator"
asm = wb.create_sheet("Assumptions")
calc = wb.create_sheet("Calc")

# ---------------- Assumptions ----------------
A = {}  # key -> absolute ref on Assumptions
asm.column_dimensions["A"].width = 34
for col in "BCD":
    asm.column_dimensions[col].width = 13
asm["A1"] = "Assumptions (edit the yellow cells to match your parts)"
asm["A1"].font = font(bold=True, size=14)
asm["A2"] = ("Values are the defaults from the web estimator. Blue text on yellow = input. "
             "A1S footprint is a placeholder until real numbers are available.")
asm["A2"].font = font(italic=True, size=9)

r = 4
header(asm, r, "Component footprints", 4)
r += 1
colheads(asm, r, ["Part", "L (mm)", "W (mm)", "H (mm)"])
asm.cell(r, 1).comment = Comment(
    "L x W is the cavity pitch on the housing, including the rib between cavities. "
    "H is cavity depth (fuses) or body height (relays).", "Estimator")
for pid, name, kind, l, w, h, _ in PARTS:
    r += 1
    asm.cell(r, 1, name).font = font()
    asm.cell(r, 1).border = BOX
    for j, v in enumerate((l, w, h)):
        inp(asm.cell(r, 2 + j), v, N1)
    A[pid] = {"L": f"Assumptions!$B${r}", "W": f"Assumptions!$C${r}", "H": f"Assumptions!$D${r}"}

r += 2
header(asm, r, "Housing geometry", 4)
for key, label, val, fmt in GEO:
    r += 1
    asm.cell(r, 1, label).font = font()
    asm.cell(r, 1).border = BOX
    inp(asm.cell(r, 2), val, fmt)
    A[key] = f"Assumptions!$B${r}"

r += 2
header(asm, r, "Resins", 4)
r += 1
colheads(asm, r, ["Resin", "Density (g/cm³)", "Price ($/kg)"])
resin_first = r + 1
for name, rho, price in RESINS:
    r += 1
    asm.cell(r, 1, name).font = font()
    asm.cell(r, 1).border = BOX
    inp(asm.cell(r, 2), rho, N2)
    inp(asm.cell(r, 3), price, MONEY2)
resin_last = r
RES_NAMES = f"Assumptions!$A${resin_first}:$A${resin_last}"
RES_RHO = f"Assumptions!$B${resin_first}:$B${resin_last}"
RES_PRICE = f"Assumptions!$C${resin_first}:$C${resin_last}"

r += 2
header(asm, r, "Molding", 4)
for key, label, val, fmt in MOLD:
    r += 1
    asm.cell(r, 1, label).font = font()
    asm.cell(r, 1).border = BOX
    inp(asm.cell(r, 2), val, fmt)
    A[key] = f"Assumptions!$B${r}"
r += 2
for line in [
    "Cycle time = base + k × wall².",
    "Clamp tonnage = projected area (in²) × tons/in² × cavities; press rate is picked from the tonnage bands.",
    "Molding cost per part = rate × cycle ÷ cavities. Material = mass × (1 + scrap) × $/kg.",
]:
    asm.cell(r, 1, line).font = font(italic=True, size=9)
    r += 1

# ---------------- Estimator (inputs) ----------------
E = {}
est.column_dimensions["A"].width = 34
for col in "BCDEFG":
    est.column_dimensions[col].width = 15
est["A1"] = "EDC Plastic Cost Estimator"
est["A1"].font = font(bold=True, size=16)
est["A2"] = "Enter fuse and relay counts. Get housing + cover size, weight and per-part cost."
est["A2"].font = font(italic=True, size=9)
est["A3"] = ("Legend: yellow cells with blue text are inputs; everything else is calculated. "
             "Material and molding assumptions live on the Assumptions tab.")
est["A3"].font = font(italic=True, size=9)

r = 5
header(est, r, "Inputs", 2)
r += 1
colheads(est, r, ["Fuses", "Count"])
for pid, name, kind, *_, cnt in PARTS:
    if kind != "fuse":
        continue
    r += 1
    est.cell(r, 1, name).font = font()
    est.cell(r, 1).border = BOX
    inp(est.cell(r, 2), cnt, N0)
    E[pid] = f"Estimator!$B${r}"
r += 1
colheads(est, r, ["PCB relays", "Count"])
for pid, name, kind, *_, cnt in PARTS:
    if kind != "relay":
        continue
    r += 1
    est.cell(r, 1, name).font = font()
    est.cell(r, 1).border = BOX
    inp(est.cell(r, 2), cnt, N0)
    E[pid] = f"Estimator!$B${r}"
count_rows = (7, r)

r += 1
colheads(est, r, ["Options", "Value"])
r += 1
est.cell(r, 1, "Relay layout (relative to fuses)").font = font()
est.cell(r, 1).border = BOX
inp(est.cell(r, 2), "Beside")
est.cell(r, 2).comment = Comment("Beside: relays add footprint area. Under: relays stack under the fuses, "
                                 "so footprint = larger of fuse or relay area.", "Estimator")
E["layout"] = f"Estimator!$B${r}"
dv_layout = DataValidation(type="list", formula1='"Beside,Under"', allow_blank=False)
est.add_data_validation(dv_layout)
dv_layout.add(f"B{r}")
r += 1
est.cell(r, 1, "Cavities per mold").font = font()
est.cell(r, 1).border = BOX
inp(est.cell(r, 2), 2, N0)
E["cav"] = f"Estimator!$B${r}"
dv_cav = DataValidation(type="list", formula1='"1,2,4,8"', allow_blank=False)
est.add_data_validation(dv_cav)
dv_cav.add(f"B{r}")
resin_dv = DataValidation(type="list", formula1=f"={RES_NAMES}", allow_blank=False)
est.add_data_validation(resin_dv)
for key, label in [("housing", "Main housing resin"), ("upper", "Upper cover resin"), ("lower", "Lower cover resin")]:
    r += 1
    est.cell(r, 1, label).font = font()
    est.cell(r, 1).border = BOX
    inp(est.cell(r, 2), "PA66-GF30")
    resin_dv.add(f"B{r}")
    E["res_" + key] = f"Estimator!$B${r}"
count_dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0",
                          error="Enter a whole number of 0 or more.", showErrorMessage=True)
est.add_data_validation(count_dv)
count_dv.add(f"B{count_rows[0]}:B{count_rows[1]}")
inputs_end = r

# ---------------- Calc ----------------
# Column B = base case; C/D/E = base + 1 of a part, for the marginal-cost line.
SCEN = [("B", "Base", None), ("C", "+1 Micro2", "micro2"), ("D", "+1 Micro3", "micro3"),
        ("E", "+1 Micro PCB relay", "relMicro")]
calc.column_dimensions["A"].width = 38
for col, *_ in SCEN:
    calc.column_dimensions[col].width = 16
calc["A1"] = "Calculation engine (no inputs here)"
calc["A1"].font = font(bold=True, size=14)
calc["A2"] = "Column B is the estimate. C-E add one part each so the Estimator can show marginal cost."
calc["A2"].font = font(italic=True, size=9)
colheads(calc, 3, ["Line item"] + [s[1] for s in SCEN])

C = {}  # name -> row
row = [3]


def line(name, label, fn, fmt=N1, section=None):
    """fn(col, ref) -> formula string for that scenario column."""
    if section:
        row[0] += 1
        calc.cell(row[0], 1, section).font = font(bold=True)
        calc.cell(row[0], 1).fill = SUB_FILL
    row[0] += 1
    rr = row[0]
    C[name] = rr
    calc.cell(rr, 1, label).font = font()
    for col, _, bump in SCEN:
        ref = lambda n, c=col: f"{c}{C[n]}"
        c = calc[f"{col}{rr}"]
        f = fn(col, ref, bump)
        c.value = f
        c.number_format = fmt
        c.font = font(GREEN if "Estimator!" in f or "Assumptions!" in f else BLACK)
    return rr


first = True
for pid, name, kind, *_ in PARTS:
    line(f"n_{pid}", f"{name} count",
         lambda col, ref, bump, p=pid: f"={E[p]}+{1 if bump == p else 0}", N0,
         section="Counts" if first else None)
    first = False

g = A
line("fuseArea", "Fuse cavity area (mm²)",
     lambda c, ref, b: "=" + "+".join(f"{ref('n_' + p)}*{g[p]['L']}*{g[p]['W']}" for p, _, k, *__ in PARTS if k == "fuse"),
     N0, section="Footprint")
line("relayArea", "Relay area (mm²)",
     lambda c, ref, b: "=" + "+".join(f"{ref('n_' + p)}*{g[p]['L']}*{g[p]['W']}" for p, _, k, *__ in PARTS if k == "relay"),
     N0)
line("maxFuseH", "Deepest fuse cavity (mm)",
     lambda c, ref, b: "=MAX(0," + ",".join(f"IF({ref('n_' + p)}>0,{g[p]['H']},0)" for p, _, k, *__ in PARTS if k == "fuse") + ")")
line("maxRelayH", "Tallest relay (mm)",
     lambda c, ref, b: "=MAX(0," + ",".join(f"IF({ref('n_' + p)}>0,{g[p]['H']},0)" for p, _, k, *__ in PARTS if k == "relay") + ")")
line("ribVol", "Cavity rib volume (mm³)",
     lambda c, ref, b: "=" + "+".join(f"{ref('n_' + p)}*({g[p]['L']}+{g[p]['W']})*{g['rib']}*{g[p]['H']}"
                                      for p, _, k, *__ in PARTS if k == "fuse"), N0)
line("comp", "Component area per layout (mm²)",
     lambda c, ref, b: f'=IF({E["layout"]}="Under",MAX({ref("fuseArea")},{ref("relayArea")}),{ref("fuseArea")}+{ref("relayArea")})', N0)
line("area", "Usable field area (mm²)",
     lambda c, ref, b: f"=IF({g['pack']}>0,{ref('comp')}/{g['pack']},0)+{g['conn']}", N0)
line("cw", "Field width (mm)", lambda c, ref, b: f"=IF({g['aspect']}>0,SQRT({ref('area')}/{g['aspect']}),0)")
line("cl", "Field length (mm)", lambda c, ref, b: f"={ref('cw')}*{g['aspect']}")

line("hL", "Housing L (mm)", lambda c, ref, b: f"={ref('cl')}+2*{g['edge']}+2*{g['wall']}", section="Part geometry")
line("hW", "Housing W (mm)", lambda c, ref, b: f"={ref('cw')}+2*{g['edge']}+2*{g['wall']}")
line("hH", "Housing H (mm)", lambda c, ref, b: f"={ref('maxFuseH')}+{g['stack']}")
line("hVol", "Housing volume (mm³)",
     lambda c, ref, b: f"={ref('hL')}*{ref('hW')}*{g['wall']}+2*({ref('hL')}+{ref('hW')})*{ref('hH')}*{g['wall']}+{ref('ribVol')}", N0)
line("cL", "Cover L (mm)", lambda c, ref, b: f"={ref('hL')}+2*({g['coverWall']}+0.5)")
line("cW", "Cover W (mm)", lambda c, ref, b: f"={ref('hW')}+2*({g['coverWall']}+0.5)")
line("uH", "Upper cover H (mm)", lambda c, ref, b: f"={g['fuseHead']}+{g['coverClear']}")
line("lH", "Lower cover H (mm)", lambda c, ref, b: f"=MAX({g['minLower']},{ref('maxRelayH')}+{g['coverClear']})")
cover_vol = lambda H: (lambda c, ref, b: f"={ref('cL')}*{ref('cW')}*{g['coverWall']}+2*({ref('cL')}+{ref('cW')})*{ref(H)}*{g['coverWall']}")
line("uVol", "Upper cover volume (mm³)", cover_vol("uH"), N0)
line("lVol", "Lower cover volume (mm³)", cover_vol("lH"), N0)

PART_DEFS = [("housing", "Main housing", "hL", "hW", "hH", g["wall"], "hVol"),
             ("upper", "Upper cover", "cL", "cW", "uH", g["coverWall"], "uVol"),
             ("lower", "Lower cover", "cL", "cW", "lH", g["coverWall"], "lVol")]
for key, title, Lk, Wk, Hk, wall, Vk in PART_DEFS:
    res = E["res_" + key]
    line(f"{key}_mass", "Mass (g)",
         lambda c, ref, b, V=Vk, res=res: f"={ref(V)}/1000*INDEX({RES_RHO},MATCH({res},{RES_NAMES},0))",
         N1, section=title)
    line(f"{key}_mat", "Material ($)",
         lambda c, ref, b, key=key, res=res: f"={ref(key + '_mass')}/1000*(1+{g['scrap']})*INDEX({RES_PRICE},MATCH({res},{RES_NAMES},0))",
         MONEY)
    line(f"{key}_tons", "Clamp tonnage (t)",
         lambda c, ref, b, L=Lk, W=Wk: f"={ref(L)}*{ref(W)}/645.16*{g['clamp']}*{E['cav']}", N0)
    line(f"{key}_cycle", "Cycle time (s)", lambda c, ref, b, wall=wall: f"={g['base']}+{g['k']}*{wall}^2")
    line(f"{key}_rate", "Press rate ($/hr)",
         lambda c, ref, b, key=key: (f"=IF({ref(key + '_tons')}<300,{g['r300']},IF({ref(key + '_tons')}<500,{g['r500']},"
                                     f"IF({ref(key + '_tons')}<800,{g['r800']},{g['rMax']})))"), MONEY2)
    line(f"{key}_mold", "Molding ($)",
         lambda c, ref, b, key=key: f"={ref(key + '_rate')}/3600*{ref(key + '_cycle')}/{E['cav']}", MONEY)
    line(f"{key}_cost", "Part cost ($)",
         lambda c, ref, b, key=key: f"={ref(key + '_mat')}+{ref(key + '_mold')}", MONEY)

line("total", "Plastic cost per EDC set ($)",
     lambda c, ref, b: f"={ref('housing_cost')}+{ref('upper_cost')}+{ref('lower_cost')}", MONEY, section="Totals")
line("mass", "Total plastic (g)",
     lambda c, ref, b: f"={ref('housing_mass')}+{ref('upper_mass')}+{ref('lower_mass')}")
line("stackH", "Stack height (mm)", lambda c, ref, b: f"={ref('hH')}+{ref('uH')}+{ref('lH')}")
for rr in (C["total"], C["mass"], C["stackH"]):
    for col in "ABCDE":
        calc[f"{col}{rr}"].font = font(bold=True)
        calc[f"{col}{rr}"].fill = TOT_FILL
calc.freeze_panes = "B4"

# ---------------- Estimator (results) ----------------
K = lambda name, col="B": f"Calc!{col}{C[name]}"
r = 5
est["D5"].value = None
for c in range(4, 8):
    est.cell(r, c).fill = HEAD_FILL
est.cell(r, 4, "Results").font = Font(name=FONT, bold=True, color="FFFFFF", size=11)


def res(row, label, formula, fmt, big=False):
    a = est.cell(row, 4, label)
    a.font = font(bold=big)
    v = est.cell(row, 5, formula)
    v.font = Font(name=FONT, color=GREEN, bold=big, size=14 if big else 10)
    v.number_format = fmt
    if big:
        a.fill = v.fill = TOT_FILL


res(6, "Plastic cost per EDC set", f"={K('total')}", MONEY2, big=True)
res(7, "Total plastic (g)", f"={K('mass')}", N0)
res(8, "Outer footprint L (mm)", f"={K('cL')}", N0)
res(9, "Outer footprint W (mm)", f"={K('cW')}", N0)
res(10, "Stack height (mm)", f"={K('stackH')}", N0)

est.cell(12, 4, "Marginal cost").font = font(bold=True)
for i, (col, label, bump) in enumerate(SCEN[1:]):
    rr = 13 + i
    est.cell(rr, 4, f"Each extra {label[3:]} adds").font = font()
    c = est.cell(rr, 5, f"=Calc!{col}{C['total']}-Calc!B{C['total']}")
    c.font = font(GREEN)
    c.number_format = MONEY

# Per-part table
r = max(inputs_end, 16) + 2
header(est, r, "Per-part breakdown", 7)
r += 1
colheads(est, r, ["", "Main housing", "Upper cover", "Lower cover"])
rows = [
    ("Resin", lambda k: f"={E['res_' + k]}", "@"),
    ("L (mm)", lambda k: f"={K({'housing': 'hL'}.get(k, 'cL'))}", N0),
    ("W (mm)", lambda k: f"={K({'housing': 'hW'}.get(k, 'cW'))}", N0),
    ("H (mm)", lambda k: f"={K({'housing': 'hH', 'upper': 'uH', 'lower': 'lH'}[k])}", N0),
    ("Volume (cm³)", lambda k: f"={K({'housing': 'hVol', 'upper': 'uVol', 'lower': 'lVol'}[k])}/1000", N1),
    ("Mass (g)", lambda k: f"={K(k + '_mass')}", N0),
    ("Material", lambda k: f"={K(k + '_mat')}", MONEY),
    ("Press tonnage (t)", lambda k: f"={K(k + '_tons')}", N0),
    ("Cycle time (s)", lambda k: f"={K(k + '_cycle')}", N0),
    ("Molding", lambda k: f"={K(k + '_mold')}", MONEY),
    ("Part cost", lambda k: f"={K(k + '_cost')}", MONEY),
]
for label, fn, fmt in rows:
    r += 1
    est.cell(r, 1, label).font = font(bold=label == "Part cost")
    est.cell(r, 1).border = BOX
    for j, key in enumerate(["housing", "upper", "lower"]):
        c = est.cell(r, 2 + j, fn(key))
        c.font = Font(name=FONT, color=GREEN, bold=label == "Part cost")
        c.number_format = fmt
        c.border = BOX
        c.alignment = Alignment(horizontal="right")
        if label == "Part cost":
            c.fill = TOT_FILL
    if label == "Part cost":
        est.cell(r, 1).fill = TOT_FILL
r += 2
est.cell(r, 1, "Not ported from the web page: the to-scale plan view and its cavity-fit warning.").font = font(italic=True, size=9)
est.freeze_panes = "A5"

for ws in (est, asm, calc):
    ws.sheet_view.showGridLines = False
wb.save(OUT)
print("wrote", OUT)
