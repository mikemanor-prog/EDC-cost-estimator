"""Recalculates EDC_Cost_Estimator_v1.1.xlsx in LibreOffice for each test case and prints the results as JSON."""
import json, os, subprocess, sys, tempfile, warnings
import openpyxl
warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "EDC_Cost_Estimator_v1.1.xlsx")
CALC_ROWS = range(5, 120)


def run(cases):
    d = tempfile.mkdtemp()
    files = []
    for i, case in enumerate(cases):
        wb = openpyxl.load_workbook(SRC)
        for k, v in case.items():
            sh, c = k.split("!")
            wb[sh][c] = v
        fn = f"{d}/c{i}.xlsx"
        wb.save(fn)
        files.append(fn)
    os.makedirs(f"{d}/o")
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", f"{d}/o"] + files,
                   capture_output=True, timeout=300)
    out = []
    for i in range(len(cases)):
        wb = openpyxl.load_workbook(f"{d}/o/c{i}.xlsx", data_only=True)
        f = openpyxl.load_workbook(files[i])["Calc"]
        ws = wb["Calc"]
        res = {}
        for r in CALC_ROWS:
            label = f.cell(r, 1).value
            if label and ws.cell(r, 2).value is not None:
                res[f"r{r}"] = {"label": label, "vals": [ws.cell(r, c).value for c in range(2, 6)]}
        out.append(res)
    return out


if __name__ == "__main__":
    print(json.dumps(run(json.load(open(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "edc_cases_v1_1.json"))))))
